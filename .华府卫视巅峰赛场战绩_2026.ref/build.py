# -*- coding: utf-8 -*-
"""生成《华府卫视巅峰赛场战绩（2026年）》xlsx。"""
import csv, datetime, os

try:
    import openpyxl
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl>=3.1.0"])
    import openpyxl

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)  # 仓库根目录（本文件所在目录的上一级）

SRC = os.path.join(ROOT, "peak-arena-ledger", "output", "巅峰赛场战绩_2026.csv")
OUT = os.path.join(ROOT, "华府卫视巅峰赛场战绩_2026.xlsx")
TITLE = "华府卫视巅峰赛场战绩（2026年）"


def xl_color(css_hex: str) -> str:
    value = css_hex.removeprefix("#").upper()
    if len(value) != 6:
        raise ValueError(f"Expected #RRGGBB, got: {css_hex}")
    return "FF" + value


XL_HEAD_BG = xl_color("#4472C4")
XL_HEAD_FG = xl_color("#FFFFFF")
XL_STRIPE = xl_color("#F2F6FC")
XL_BORDER = xl_color("#BFBFBF")
XL_NOTE_BG = xl_color("#EDF2FB")

thin = Side(style="thin", color=XL_BORDER)
border_bottom = Border(bottom=thin)

rows = []
with open(SRC, encoding="utf-8-sig") as f:
    for r in csv.DictReader(f):
        rows.append(r)

wb = Workbook()
wb.properties.title = TITLE

# ---------- Sheet 1: 战绩明细 ----------
ws = wb.active
ws.title = "巅峰赛场战绩"

# 锚点：第1行标题、第2行表头、第3~206行数据
HEAD_ROW = 2
FIRST = 3
LAST = FIRST + len(rows) - 1

ws["A1"] = TITLE
ws["A1"].font = Font(size=14, bold=True, color=xl_color("#1F3864"))
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)
ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
ws.row_dimensions[1].height = 26

headers = ["序号", "日期", "选手1", "比分", "选手2"]
for c, h in enumerate(headers, 1):
    cell = ws.cell(row=HEAD_ROW, column=c, value=h)
    cell.font = Font(bold=True, color=XL_HEAD_FG)
    cell.fill = PatternFill("solid", fgColor=XL_HEAD_BG)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = border_bottom
ws.row_dimensions[HEAD_ROW].height = 20

for i, r in enumerate(rows):
    rr = FIRST + i
    d = datetime.datetime.strptime(r["日期"], "%Y-%m-%d").date()
    ws.cell(row=rr, column=1, value=i + 1).alignment = Alignment(horizontal="center")
    c2 = ws.cell(row=rr, column=2, value=d)
    c2.number_format = "yyyy-mm-dd"
    c2.alignment = Alignment(horizontal="center")
    ws.cell(row=rr, column=3, value=r["选手1"]).alignment = Alignment(horizontal="left")
    sc = ws.cell(row=rr, column=4, value=r["比分"])
    sc.alignment = Alignment(horizontal="center")
    sc.font = Font(bold=True, color=xl_color("#2563EB"))
    ws.cell(row=rr, column=5, value=r["选手2"]).alignment = Alignment(horizontal="left")
    if i % 2 == 1:
        for c in range(1, 6):
            ws.cell(row=rr, column=c).fill = PatternFill("solid", fgColor=XL_STRIPE)

for col, w in zip("ABCDE", [6, 13, 14, 9, 14]):
    ws.column_dimensions[col].width = w

ws.freeze_panes = "A3"
ws.auto_filter.ref = "A%d:E%d" % (HEAD_ROW, LAST)

# ---------- Sheet 2: 数据说明 ----------
ws2 = wb.create_sheet("数据说明")
ws2.column_dimensions["A"].width = 16
ws2.column_dimensions["B"].width = 92
ws2["A1"] = "数据说明"
ws2["A1"].font = Font(size=13, bold=True, color=xl_color("#1F3864"))
ws2.merge_cells("A1:B1")
notes = [
    ("数据来源", "微信公众号「鸡坛快迅」发布的「华府卫视昨日战报」，经 wxredian.com 归档逐篇抓取原文。"),
    ("统计范围", "2026-01-01 至 2026-09-18，共 %d 场「巅峰赛场」对局。" % len(rows)),
    ("日期口径", "战报为次日发布，标题即「昨日战报」，故表中日期 = 战报发布日 − 1 天，即实际比赛日。"),
    ("抓取范围", "该公众号 2025-12-01 至 2026-09-20 全部 277 篇推文，其中 196 篇含「巅峰赛场」条目。"),
    ("说明", "同一日出现两行（如 bo7+bo3）时，为该日安排了两场巅峰赛场对局。"),
    ("抓取时间", "2026-09-21"),
]
for i, (k, v) in enumerate(notes):
    r = 3 + i
    a = ws2.cell(row=r, column=1, value=k)
    a.font = Font(bold=True)
    a.alignment = Alignment(vertical="top")
    a.fill = PatternFill("solid", fgColor=XL_NOTE_BG)
    b = ws2.cell(row=r, column=2, value=v)
    b.alignment = Alignment(vertical="top", wrap_text=True)

wb.save(OUT)
print("saved:", OUT, "rows:", len(rows))
